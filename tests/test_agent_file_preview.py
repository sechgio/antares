from __future__ import annotations

import base64
import io
import struct
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from backend.core.agent import file_preview


def test_pdf_pages_are_rendered_as_bounded_images_without_modifying_source(tmp_path: Path):
    import pymupdf
    from PIL import Image

    path = tmp_path / "informe.pdf"
    with pymupdf.open() as document:
        document.new_page(width=2400, height=1800).insert_text((72, 72), "Primera")
        document.new_page().insert_text((72, 72), "Segunda")
        document.save(path)
    original = path.read_bytes()
    result = file_preview.render_file(path, {"page": 1})
    assert result["kind"] == "pdf" and result["page"] == 1 and result["pages"] == 2
    assert result["name"] == path.name and result["size"] == len(original)
    with Image.open(io.BytesIO(base64.b64decode(result["png_base64"]))) as image:
        assert max(image.size) <= 1600
    first = file_preview.render_file(path, {})
    with Image.open(io.BytesIO(base64.b64decode(first["png_base64"]))) as image:
        assert max(image.size) <= 1600
    assert path.read_bytes() == original
    with pytest.raises(ValueError, match="página"):
        file_preview.render_file(path, {"page": 2})


def test_encrypted_pdf_has_an_actionable_error(tmp_path: Path):
    import pymupdf

    path = tmp_path / "protegido.pdf"
    with pymupdf.open() as document:
        document.new_page()
        document.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="reader")
    with pytest.raises(ValueError, match="contraseña"):
        file_preview.render_file(path, {})


def test_excel_preserves_sheet_and_row_order_with_pagination(tmp_path: Path):
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "datos.xlsx"
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Resumen"
    first.append(["Código", "Importe", "Fórmula"])
    first.append(["001", 12, "=2+2"])
    second = workbook.create_sheet("Detalle")
    for number in range(55):
        second.append([number, f"Fila {number}"])
    workbook.save(path)
    original = path.read_bytes()
    result = file_preview.render_file(path, {"sheet": 1, "page": 1})
    assert result["sheets"] == ["Resumen", "Detalle"] and result["sheet"] == 1
    assert result["page"] == 1 and result["pages"] == 2 and result["row_start"] == 51
    assert result["total_rows"] == 55 and result["total_columns"] == 2
    assert result["rows"] == [[str(number), f"Fila {number}"] for number in range(50, 55)]
    assert result["truncated"] is False
    assert file_preview.render_file(path, {})["rows"] == [["Código", "Importe", "Fórmula"], ["001", "12", ""]]
    assert path.read_bytes() == original
    reloaded = load_workbook(path)
    assert reloaded["Resumen"]["C2"].value == "=2+2"
    reloaded.close()
    for params in ({"sheet": 2}, {"sheet": 1, "page": 2}):
        with pytest.raises(ValueError, match="existe"):
            file_preview.render_file(path, params)


def test_excel_bounds_wide_and_long_cells(tmp_path: Path):
    from openpyxl import Workbook

    path = tmp_path / "ancho.xlsx"
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.append(["a" * 600, *range(55)])
    workbook.save(path)
    result = file_preview.render_file(path, {})
    assert result["total_columns"] == 56 and len(result["rows"][0]) == 50
    assert len(result["rows"][0][0]) == 500 and result["truncated"]


def test_legacy_excel_is_read_without_contaminating_json_rpc_stdout(tmp_path: Path, capsys: pytest.CaptureFixture[str]):
    def record(opcode: int, data: bytes) -> bytes:
        return struct.pack("<HH", opcode, len(data)) + data

    path = tmp_path / "antiguo.xls"
    path.write_bytes(record(0x0409, struct.pack("<HH", 0x0400, 0x0010))
                     + record(0x0203, struct.pack("<HHHd", 0, 0, 0, 12.5)) + record(0x000A, b""))
    result = file_preview.render_file(path, {})
    assert result["kind"] == "spreadsheet" and result["rows"] == [["12.5"]]
    assert result["pages"] == 1 and result["sheet"] == 0
    assert capsys.readouterr().out == ""


@pytest.mark.parametrize("encoding,delimiter,suffix", [("utf-8-sig", ",", ".csv"), ("cp1252", ";", ".csv"),
                                                      ("utf-16", "\t", ".tsv")])
def test_csv_preserves_unicode_quoted_delimiters_and_newlines(tmp_path: Path, encoding: str, delimiter: str, suffix: str):
    path = tmp_path / f"registros{suffix}"
    path.write_bytes(f'Nombre{delimiter}Detalle\n"Peña"{delimiter}"Uno{delimiter}dos\ntres"\n'.encode(encoding))
    result = file_preview.render_file(path, {})
    assert result["kind"] == "spreadsheet" and result["rows"] == [["Nombre", "Detalle"], ["Peña", f"Uno{delimiter}dos\ntres"]]
    assert result["total_rows"] == 2 and result["pages"] == 1
    with pytest.raises(ValueError, match="hoja"):
        file_preview.render_file(path, {"sheet": 1})


def test_csv_row_limit_and_pagination(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    path = tmp_path / "datos.csv"
    path.write_text("\n".join(f"{number},valor" for number in range(55)), encoding="utf-8")
    result = file_preview.render_file(path, {"page": 1})
    assert result["row_start"] == 51 and len(result["rows"]) == 5
    monkeypatch.setattr(file_preview, "MAX_ROWS", 51)
    limited = file_preview.render_file(path, {"page": 1})
    assert limited["truncated"] and limited["total_rows"] == 51 and len(limited["rows"]) == 1


def test_docx_preserves_paragraph_table_order_and_exposes_no_active_html(tmp_path: Path):
    from docx import Document

    path = tmp_path / "carta.docx"
    document = Document()
    document.add_heading("Título", level=1)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Código"
    table.cell(0, 1).text = "<script>alert(1)</script>"
    document.add_paragraph("Texto final")
    document.save(path)
    original = path.read_bytes()
    result = file_preview.render_file(path, {})
    assert result["kind"] == "document" and result["truncated"] is False
    assert result["blocks"] == [
        {"kind": "paragraph", "text": "Título", "style": "h1"},
        {"kind": "table", "rows": [["Código", "<script>alert(1)</script>"]]},
        {"kind": "paragraph", "text": "Texto final", "style": "normal"},
    ]
    assert path.read_bytes() == original


def test_docx_content_is_bounded(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from docx import Document

    path = tmp_path / "largo.docx"
    document = Document()
    document.add_paragraph("Texto muy largo")
    document.add_paragraph("Siguiente")
    document.save(path)
    monkeypatch.setattr(file_preview, "MAX_TEXT_CHARS", 8)
    result = file_preview.render_file(path, {})
    assert result["truncated"] and result["blocks"] == [{"kind": "paragraph", "text": "Texto mu", "style": "normal"}]


def test_image_is_rasterized_and_scaled(tmp_path: Path):
    from PIL import Image

    path = tmp_path / "foto.jpg"
    Image.new("RGB", (3000, 1800), "red").save(path)
    original = path.read_bytes()
    result = file_preview.render_file(path, {})
    assert result["kind"] == "image" and result["width"] == 3000 and result["height"] == 1800
    with Image.open(io.BytesIO(base64.b64decode(result["png_base64"]))) as shown:
        assert shown.size == (1600, 960) and shown.format == "PNG"
    assert path.read_bytes() == original


def test_text_and_html_are_returned_as_inert_bounded_text(tmp_path: Path):
    path = tmp_path / "pagina.html"
    path.write_bytes("<script>alert('hola')</script>\nPeña".encode())
    result = file_preview.render_file(path, {})
    assert result["kind"] == "text" and result["text"] == path.read_text(encoding="utf-8")
    path = tmp_path / "largo.txt"
    path.write_text("á" * 200_001, encoding="utf-8")
    result = file_preview.render_file(path, {})
    assert result["text"] == "á" * 200_000 and result["truncated"]


def test_text_decodes_bom_and_windows_encoding(tmp_path: Path):
    path = tmp_path / "texto.txt"
    for encoding in ("utf-16", "cp1252", "utf-8-sig"):
        path.write_bytes("Peña\nñ".encode(encoding))
        assert file_preview.render_file(path, {})["text"] == "Peña\nñ"


@pytest.mark.parametrize("suffix", [".pdf", ".xlsx", ".docx", ".xls", ".jpg"])
def test_corrupt_files_raise_readable_errors(tmp_path: Path, suffix: str):
    path = tmp_path / f"roto{suffix}"
    path.write_bytes(b"this is not a document")
    with pytest.raises(ValueError, match="leer"):
        file_preview.render_file(path, {})


def test_binary_unsupported_missing_directory_and_large_files(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    path = tmp_path / "binario.txt"
    path.write_bytes(b"MZ\x00\x00")
    with pytest.raises(ValueError, match="binario"):
        file_preview.render_file(path, {})
    unsupported = tmp_path / "programa.exe"
    unsupported.write_bytes(b"MZ")
    with pytest.raises(ValueError, match="formato"):
        file_preview.render_file(unsupported, {})
    for missing in (tmp_path / "ausente.txt", tmp_path):
        with pytest.raises(ValueError, match="disponible"):
            file_preview.render_file(missing, {})
    monkeypatch.setattr(file_preview, "MAX_FILE_BYTES", 1)
    with pytest.raises(ValueError, match="50 MiB"):
        file_preview.render_file(path, {})


@pytest.mark.parametrize("params", [{"page": True}, {"page": -1}, {"page": "1"}, {"sheet": 1.5}, {"sheet": False}])
def test_preview_rejects_invalid_page_and_sheet(tmp_path: Path, params: dict):
    path = tmp_path / "texto.txt"
    path.write_text("dato")
    with pytest.raises(ValueError, match="índices"):
        file_preview.render_file(path, params)


def test_archive_size_is_checked_before_opening_document(tmp_path: Path):
    path = tmp_path / "comprimido.docx"
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (32 * 1024 * 1024 + 1))
    with pytest.raises(ValueError, match="32 MiB"):
        file_preview.render_file(path, {})


def test_archive_entry_count_is_bounded(tmp_path: Path):
    path = tmp_path / "muchos.xlsx"
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for index in range(4097):
            archive.writestr(f"entry-{index}", "")
    with pytest.raises(ValueError, match="descomprimido"):
        file_preview.render_file(path, {})

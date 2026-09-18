from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from backend.core.fichas_tecnicas.importer import (
    import_fichas_from_bytes,
    parse_csv_bytes,
    parse_xlsx_bytes,
)


def _xlsx_bytes(rows: list[list[object]], number_formats: dict[tuple[int, int], str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    if number_formats:
        for (r, c), fmt in number_formats.items():
            ws.cell(row=r, column=c).number_format = fmt
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_csv_prefiere_semicolon_y_normaliza_headers() -> None:
    content = b"Cliente;Distrito\nACME;Lima\n"
    rows = parse_csv_bytes(content)
    assert rows == [{"cliente": "ACME", "distrito": "Lima"}]


def test_parse_csv_cae_a_latin1_cuando_utf8_falla() -> None:
    content = "Cliente\nJosé\n".encode("latin-1")
    rows = parse_csv_bytes(content)
    assert rows[0]["cliente"] == "José"


def test_parse_csv_descarta_filas_vacias() -> None:
    content = b"a;b\n1;2\n;\n"
    rows = parse_csv_bytes(content)
    assert rows == [{"a": "1", "b": "2"}]


def test_parse_csv_bytes_arbitrarios_decodifican_via_latin1() -> None:
    # latin-1 mapea los 256 bytes: la rama "indecodificable" es inalcanzable
    rows = parse_csv_bytes(b"a,b\n\xff,\xfe\n")
    assert len(rows) == 1


def test_parse_xlsx_normaliza_y_omite_vacias() -> None:
    content = _xlsx_bytes(
        [
            ["Cliente", "Distrito"],
            ["ACME", "Lima"],
            [None, None],
            ["", ""],
            ["XYZ", "Surco"],
        ]
    )
    rows = parse_xlsx_bytes(content)
    assert rows == [
        {"cliente": "ACME", "distrito": "Lima"},
        {"cliente": "XYZ", "distrito": "Surco"},
    ]


def test_parse_xlsx_concentracion_respeta_number_format() -> None:
    content = _xlsx_bytes(
        [["producto_1_concentracion"], [0.5]],
        number_formats={(2, 1): "0.000"},
    )
    rows = parse_xlsx_bytes(content)
    assert rows[0]["producto_1_concentracion"] == "0.500"


def test_parse_xlsx_sin_filas_devuelve_vacio() -> None:
    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    # hoja activa sin datos → iter_rows produce una celda vacía; headers + nada útil
    assert parse_xlsx_bytes(buf.getvalue()) == []


def test_import_fichas_rechaza_extension_desconocida_y_vacio() -> None:
    with pytest.raises(ValueError, match="no soportado"):
        import_fichas_from_bytes("datos.txt", b"a,b\n1,2\n")
    with pytest.raises(ValueError, match="vacio"):
        import_fichas_from_bytes("datos.csv", b"a;b\n;\n")


def test_import_fichas_csv_genera_fichas_normalizadas() -> None:
    csv_data = b"cliente;distrito;servicio_desinfeccion\nACME;Lima;SI\nXYZ;Surco;no\n"
    fichas = import_fichas_from_bytes("fichas.csv", csv_data)
    assert len(fichas) == 2
    assert fichas[0]["cliente"] == "ACME"
    assert fichas[0]["servicio"]["desinfeccion"] is True
    assert fichas[1]["servicio"]["desinfeccion"] is False
    assert fichas[0]["status"] == "draft"
    assert fichas[0]["id"] != fichas[1]["id"]


def test_import_fichas_xlsx_end_to_end() -> None:
    content = _xlsx_bytes([["cliente", "hora_inicio"], ["ACME", "08:00"]])
    fichas = import_fichas_from_bytes("fichas.xlsx", content)
    assert fichas[0]["cliente"] == "ACME"
    assert fichas[0]["hora_inicio"] == "08:00"

from backend.core import tabular_report_import
from backend.core.tabular_report_import import (
    normalize_tabular_header,
    normalize_tabular_key,
    parse_csv_rows,
)


def test_normalize_tabular_header_preserves_domain_quote_policy() -> None:
    assert normalize_tabular_header('Valv Cond 2"', strip_quotes=True) == "valvcond2"
    assert normalize_tabular_header('Valv Cond 2"', strip_quotes=False) == 'valvcond2"'


def test_normalize_tabular_key_uses_mapping_then_falls_back_to_snake_case() -> None:
    mapping = {"fechaejecucion": "fecha_ejecucion"}

    def normalize_header(value: str) -> str:
        return normalize_tabular_header(value, strip_quotes=True)

    assert normalize_tabular_key("Fecha de Ejecución", mapping, normalize_header) == "fecha_ejecucion"
    assert normalize_tabular_key("Dato Libre", mapping, normalize_header) == "dato_libre"


def test_parse_csv_rows_lee_cp1252_antes_que_latin1() -> None:
    # En cp1252 0x97 y 0x92 son raya y comilla tipográfica; latin-1 las lee como control.
    content = "Observacion;Detalle\nZona 2 \u2014 alta\u2019 ok;listo\n".encode("cp1252")

    rows = parse_csv_rows(content, column_mapping={}, normalize_key=lambda value: str(value))

    assert rows[0]["Observacion"] == "Zona 2 \u2014 alta\u2019 ok"


def test_parse_csv_rows_solo_materializa_el_delimitador_elegido(monkeypatch) -> None:
    content = "fecha;valor\n" + "".join(f"{index};v{index}\n" for index in range(40))
    original_reader = tabular_report_import.csv.DictReader
    yielded_rows = 0

    class CountingDictReader:
        def __init__(self, *args, **kwargs):
            self.reader = original_reader(*args, **kwargs)

        @property
        def fieldnames(self):
            return self.reader.fieldnames

        def __iter__(self):
            return self

        def __next__(self):
            nonlocal yielded_rows
            row = next(self.reader)
            yielded_rows += 1
            return row

    monkeypatch.setattr(tabular_report_import.csv, "DictReader", CountingDictReader)

    rows = parse_csv_rows(
        content.encode(),
        column_mapping={"fecha": "fecha", "valor": "valor"},
        normalize_key=str.lower,
    )

    assert len(rows) == 40
    assert rows[0] == {"fecha": "0", "valor": "v0"}
    assert yielded_rows <= 43


def test_parse_xlsx_rows_usa_lectura_read_only_y_selecciona_la_mejor_hoja(monkeypatch) -> None:
    from io import BytesIO

    from openpyxl import Workbook, load_workbook

    workbook = Workbook()
    first = workbook.active
    first.append(["fecha"])
    first.append(["1"])
    second = workbook.create_sheet("mejor")
    second.append(["fecha", "valor"])
    second.append(["2", "ok"])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    def load_read_only(source, **kwargs):
        assert kwargs.get("read_only") is True
        return load_workbook(source, **kwargs)

    monkeypatch.setattr(tabular_report_import, "load_workbook", load_read_only)
    rows = tabular_report_import.parse_xlsx_rows(
        buffer.getvalue(),
        column_mapping={"fecha": "fecha", "valor": "valor"},
        normalize_header=str.lower,
        normalize_key=str.lower,
    )

    assert rows == [{"fecha": "2", "valor": "ok"}]

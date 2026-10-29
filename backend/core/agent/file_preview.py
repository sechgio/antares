from __future__ import annotations

import base64
import codecs
import csv
import io
import math
from importlib import import_module
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from backend.core.integrations.types import JsonObject

MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_TEXT_CHARS = 200_000
MAX_ROWS = 100_000
PAGE_ROWS = 50
MAX_COLUMNS = 50
MAX_CELL_CHARS = 500
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
_TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".json", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf",
    ".log", ".xml", ".html", ".htm", ".svg", ".css", ".scss", ".js", ".jsx", ".mjs", ".cjs", ".ts", ".tsx",
    ".py", ".sql", ".sh", ".ps1", ".bat", ".rs", ".go", ".java", ".c", ".h", ".cpp", ".cs", ".rb",
}


def _index(params: JsonObject, name: str) -> int:
    value = params.get(name, 0)
    if not isinstance(value, int) or isinstance(value, bool) or value < 0:
        raise ValueError("La página y la hoja deben ser índices enteros no negativos")
    return value


def _check_zip(path: Path) -> None:
    with ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 4096 or sum(entry.file_size for entry in entries) > 100 * 1024 * 1024:
            raise ValueError("El documento excede el límite de contenido descomprimido")
        if any(entry.file_size > 32 * 1024 * 1024 for entry in entries):
            raise ValueError("Una parte del documento excede 32 MiB descomprimidos")
        if any(entry.flag_bits & 1 for entry in entries):
            raise ValueError("No se pueden previsualizar documentos cifrados")


def _pdf(path: Path, page: int) -> JsonObject:
    import pymupdf

    try:
        document = pymupdf.open(path)
    except pymupdf.FileDataError as err:
        raise ValueError("No se pudo leer el PDF; puede estar dañado") from err
    with document:
        if document.needs_pass:
            raise ValueError("El PDF está protegido con contraseña")
        if page >= len(document):
            raise ValueError("La página no existe en este PDF")
        selected = document[page]
        dimension = max(selected.rect.width, selected.rect.height)
        if not math.isfinite(dimension) or dimension <= 0:
            raise ValueError("Las dimensiones de la página no son válidas")
        scale = min(2, 1600 / dimension)
        png = selected.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False).tobytes("png")
        return {"kind": "pdf", "png_base64": base64.b64encode(png).decode("ascii"),
                "page": page, "pages": len(document)}


def _image(path: Path) -> JsonObject:
    from PIL import Image, ImageOps

    try:
        with Image.open(path) as image:
            width, height = image.size
            if width * height > 40_000_000:
                raise ValueError("La imagen excede 40 millones de píxeles")
            image.thumbnail((1600, 1600))
            shown = ImageOps.exif_transpose(image).convert("RGBA")
            output = io.BytesIO()
            shown.save(output, format="PNG")
    except Image.DecompressionBombError as err:
        raise ValueError("La imagen excede el límite de píxeles") from err
    return {"kind": "image", "png_base64": base64.b64encode(output.getvalue()).decode("ascii"),
            "width": width, "height": height}


def _cell(value: object) -> str:
    return "" if value is None else str(value)[:MAX_CELL_CHARS]


def _sheet_result(sheets: list[str], sheet: int, page: int, total_rows: int, total_columns: int,
                  rows: list[list[str]], truncated: bool) -> JsonObject:
    pages = max(1, math.ceil(min(total_rows, MAX_ROWS) / PAGE_ROWS))
    if page >= pages:
        raise ValueError("La página no existe en esta hoja")
    return {"kind": "spreadsheet", "sheets": sheets, "sheet": sheet, "page": page, "pages": pages,
            "total_rows": total_rows, "total_columns": total_columns, "rows": rows,
            "row_start": page * PAGE_ROWS + 1,
            "truncated": truncated or total_rows > MAX_ROWS or total_columns > MAX_COLUMNS}


def _xlsx(path: Path, sheet: int, page: int) -> JsonObject:
    from openpyxl import load_workbook

    _check_zip(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except ValueError as err:
        raise ValueError("No se pudo leer el libro Excel; puede estar dañado") from err
    try:
        sheets = workbook.sheetnames
        if sheet >= len(sheets):
            raise ValueError("La hoja no existe en este archivo")
        selected = workbook[sheets[sheet]]
        total_rows, total_columns = selected.max_row or 0, selected.max_column or 0
        if page >= max(1, math.ceil(min(total_rows, MAX_ROWS) / PAGE_ROWS)):
            raise ValueError("La página no existe en esta hoja")
        values = selected.iter_rows(min_row=page * PAGE_ROWS + 1, max_row=min((page + 1) * PAGE_ROWS, total_rows),
                                    max_col=min(total_columns, MAX_COLUMNS), values_only=True) if total_rows else []
        rows: list[list[str]] = []
        truncated = False
        for row in values:
            truncated = truncated or any(len(str(value)) > MAX_CELL_CHARS for value in row if value is not None)
            rows.append([_cell(value) for value in row])
        return _sheet_result(sheets, sheet, page, total_rows, total_columns, rows, truncated)
    finally:
        workbook.close()


def _xls(path: Path, sheet: int, page: int) -> JsonObject:
    xlrd = import_module("xlrd")
    compdoc = import_module("xlrd.compdoc")

    try:
        workbook = xlrd.open_workbook(str(path), on_demand=True, logfile=io.StringIO())
    except (xlrd.XLRDError, compdoc.CompDocError) as err:
        raise ValueError("No se pudo leer el libro Excel; puede estar dañado o cifrado") from err
    try:
        sheets = workbook.sheet_names()
        if sheet >= len(sheets):
            raise ValueError("La hoja no existe en este archivo")
        selected = workbook.sheet_by_index(sheet)
        rows: list[list[str]] = []
        truncated = False
        for index in range(page * PAGE_ROWS, min((page + 1) * PAGE_ROWS, selected.nrows, MAX_ROWS)):
            values = selected.row_values(index, 0, min(selected.ncols, MAX_COLUMNS))
            truncated = truncated or any(len(str(value)) > MAX_CELL_CHARS for value in values)
            rows.append([_cell(value) for value in values])
        return _sheet_result(sheets, sheet, page, selected.nrows, selected.ncols, rows, truncated)
    finally:
        workbook.release_resources()


def _encoding(sample: bytes) -> str:
    if sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16"
    try:
        codecs.getincrementaldecoder("utf-8-sig")().decode(sample, final=False)
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "cp1252"


def _csv(path: Path, sheet: int, page: int) -> JsonObject:
    if sheet:
        raise ValueError("El archivo de texto solo tiene una hoja")
    with path.open("rb") as source:
        sample = source.read(65_536)
    encoding = _encoding(sample)
    decoded = sample.decode(encoding, errors="replace")
    try:
        delimiter = csv.Sniffer().sniff(decoded, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    rows: list[list[str]] = []
    total_rows = total_columns = 0
    truncated = False
    with path.open("r", encoding=encoding, errors="replace", newline="") as source:
        for index, row in enumerate(csv.reader(source, delimiter=delimiter)):
            if index >= MAX_ROWS:
                truncated = True
                break
            total_rows += 1
            total_columns = max(total_columns, len(row))
            if page * PAGE_ROWS <= index < (page + 1) * PAGE_ROWS:
                truncated = truncated or any(len(value) > MAX_CELL_CHARS for value in row[:MAX_COLUMNS])
                rows.append([_cell(value) for value in row[:MAX_COLUMNS]])
    return _sheet_result([path.stem], sheet, page, total_rows, total_columns, rows, truncated)


def _docx(path: Path) -> JsonObject:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_zip(path)
    try:
        document = Document(str(path))
    except ValueError as err:
        raise ValueError("No se pudo leer el documento Word; puede estar dañado") from err
    blocks: list[JsonObject] = []
    remaining = MAX_TEXT_CHARS
    truncated = False
    for element in document.element.body:
        if remaining <= 0 or len(blocks) >= 2000:
            truncated = True
            break
        if element.tag.endswith("}p"):
            paragraph = Paragraph(element, document)
            text = paragraph.text
            style_name = paragraph.style.name if paragraph.style is not None else ""
            style = {"Heading 1": "h1", "Heading 2": "h2", "Heading 3": "h3", "Title": "h1"}.get(style_name, "normal")
            blocks.append({"kind": "paragraph", "text": text[:remaining], "style": style})
            truncated = truncated or len(text) > remaining
            remaining -= min(len(text), remaining)
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            rows: list[list[str]] = []
            for row_index, row in enumerate(table.rows):
                if row_index >= 500 or remaining <= 0:
                    truncated = True
                    break
                cells: list[str] = []
                for cell in row.cells[:MAX_COLUMNS]:
                    text = cell.text
                    limit = min(2000, remaining)
                    cells.append(text[:limit])
                    truncated = truncated or len(text) > limit
                    remaining -= min(len(text), limit)
                truncated = truncated or len(row.cells) > MAX_COLUMNS
                rows.append(cells)
            blocks.append({"kind": "table", "rows": rows})
    return {"kind": "document", "blocks": blocks, "truncated": truncated}


def _text(path: Path) -> JsonObject:
    with path.open("rb") as source:
        data = source.read(MAX_TEXT_CHARS * 4 + 4)
    encoding = _encoding(data[:65_536])
    if encoding != "utf-16" and b"\x00" in data:
        raise ValueError("El archivo es binario y no tiene vista previa disponible")
    text = data.decode(encoding, errors="replace")
    return {"kind": "text", "text": text[:MAX_TEXT_CHARS], "truncated": len(text) > MAX_TEXT_CHARS}


def render_file(path: Path, params: JsonObject) -> JsonObject:
    page, sheet = _index(params, "page"), _index(params, "sheet")
    if not path.is_file():
        raise ValueError("El archivo ya no está disponible")
    size = path.stat().st_size
    if size > MAX_FILE_BYTES:
        raise ValueError("La vista previa admite archivos de hasta 50 MiB")
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            result = _pdf(path, page)
        elif suffix in _IMAGE_SUFFIXES:
            result = _image(path)
        elif suffix in {".xlsx", ".xlsm"}:
            result = _xlsx(path, sheet, page)
        elif suffix == ".xls":
            result = _xls(path, sheet, page)
        elif suffix in {".csv", ".tsv"}:
            result = _csv(path, sheet, page)
        elif suffix == ".docx":
            result = _docx(path)
        elif suffix in _TEXT_SUFFIXES or (not suffix and path.name.lower() in {"readme", "license", "makefile", "dockerfile"}):
            result = _text(path)
        else:
            raise ValueError("Este formato no tiene vista previa; usa PDF, Excel, Word (.docx), imágenes o texto")
    except (BadZipFile, OSError, csv.Error, KeyError, ParseError, SyntaxError) as err:
        raise ValueError("No se pudo leer el archivo; puede estar dañado, cifrado o no disponible") from err
    return {"name": path.name, "size": size, **result}

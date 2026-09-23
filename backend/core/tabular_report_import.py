from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections.abc import Callable, Mapping
from typing import Any

from openpyxl import load_workbook

ReportRow = dict[str, Any]  # allowlist: tabular imports accept heterogeneous spreadsheet values
NormalizeValue = Callable[[str], str]


def normalize_tabular_header(value: str, *, strip_quotes: bool) -> str:
    if not value:
        return ""
    text = str(value).strip().lower().replace("\ufeff", "")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = re.sub(r"\b(de|del|la|el)\b", "", text)
    separators = r'''[\s_\.:\-°/()"']+''' if strip_quotes else r"[\s_\.:\-°/()]+"
    return re.sub(separators, "", text)


def normalize_tabular_key(
    value: str,
    column_mapping: Mapping[str, str],
    normalize_header: NormalizeValue,
) -> str:
    mapped = column_mapping.get(normalize_header(value))
    if mapped:
        return mapped
    text = str(value or "").strip().lower().replace("\ufeff", "")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def import_reports(
    filename: str,
    content: bytes,
    *,
    parse_csv: Callable[[bytes], list[ReportRow]],
    parse_xlsx: Callable[[bytes], list[ReportRow]],
    transform_row: Callable[[ReportRow, int], ReportRow],
    to_int: Callable[[Any, int], int],
) -> list[ReportRow]:
    lower_name = filename.lower()
    if lower_name.endswith(".csv"):
        rows = parse_csv(content)
    elif lower_name.endswith((".xlsx", ".xls")):
        rows = parse_xlsx(content)
    else:
        msg = "Formato no soportado. Use archivos .csv o .xlsx"
        raise ValueError(msg)
    if not rows:
        msg = "El archivo esta vacio o no tiene datos validos"
        raise ValueError(msg)

    reports: list[ReportRow] = []
    used_numbers: set[int] = set()
    next_report_number = 1

    for row in rows:
        explicit_number = to_int(row.get("informe_id"), 0)
        if explicit_number > 0 and explicit_number not in used_numbers:
            report_number = explicit_number
        else:
            while next_report_number in used_numbers:
                next_report_number += 1
            report_number = next_report_number

        used_numbers.add(report_number)
        next_report_number = max(next_report_number, report_number) + 1
        reports.append(transform_row(row, report_number))

    return reports


def parse_csv_rows(
    content: bytes,
    *,
    column_mapping: Mapping[str, str],
    normalize_key: NormalizeValue,
) -> list[ReportRow]:
    text: str | None = None
    # cp1252 va antes que latin-1: latin-1 decodifica cualquier byte y se quedaría
    # con los caracteres de control en lugar de la raya y las comillas tipográficas.
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1", "iso-8859-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if text is None:
        msg = "No se pudo decodificar el archivo CSV"
        raise ValueError(msg)

    def normalize_row(row: dict[str | None, Any]) -> ReportRow:
        return {normalize_key(key): value for key, value in row.items() if key is not None}

    def read_rows(delimiter: str) -> list[ReportRow]:
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        parsed_rows: list[ReportRow] = []
        for row in reader:
            normalized = normalize_row(row)
            if any(str(value or "").strip() for value in normalized.values()):
                parsed_rows.append(normalized)
        return parsed_rows

    mapped_targets = set(column_mapping.values())
    candidates: dict[str, tuple[int, bool]] = {}
    for delimiter in (";", ",", "\t"):
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        header_keys = {normalize_key(key) for key in reader.fieldnames or []}
        mapped_column_count = sum(key in mapped_targets for key in header_keys if key)
        has_rows = False
        for row in reader:
            normalized = normalize_row(row)
            if any(str(value or "").strip() for value in normalized.values()):
                has_rows = True
                break
        candidates[delimiter] = (mapped_column_count, has_rows)

    best_column_count = max(score for score, _ in candidates.values())
    if best_column_count >= 1:
        for preferred_delimiter in (";", "\t", ","):
            score, has_rows = candidates[preferred_delimiter]
            if score == best_column_count and has_rows:
                return read_rows(preferred_delimiter)

    fallback_delimiter = next(
        (delimiter for delimiter in (";", ",", "\t") if candidates[delimiter][1]),
        None,
    )
    return read_rows(fallback_delimiter) if fallback_delimiter else []


def parse_xlsx_rows(
    content: bytes,
    *,
    column_mapping: Mapping[str, str],
    normalize_header: NormalizeValue,
    normalize_key: NormalizeValue,
) -> list[ReportRow]:
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as err:
        msg = f"Error al abrir el archivo Excel: {err}"
        raise ValueError(msg) from err

    try:
        selected_sheet = None
        selected_header_index = -1
        selected_headers: tuple[Any, ...] = ()
        best_sheet_score = -1

        for worksheet in workbook.worksheets:
            worksheet_headers = list(worksheet.iter_rows(min_row=1, max_row=10, values_only=True))
            if not worksheet_headers:
                continue

            header_index = -1
            best_score = -1
            for index, values in enumerate(worksheet_headers):
                non_empty = [value for value in values if value is not None and str(value).strip() != ""]
                if not non_empty:
                    continue
                score = sum(normalize_header(str(value)) in column_mapping for value in non_empty)
                if score > best_score:
                    best_score = score
                    header_index = index

            if header_index == -1 or best_score <= 0:
                for index, values in enumerate(worksheet_headers):
                    if any(value is not None and str(value).strip() != "" for value in values):
                        header_index = index
                        break

            if header_index == -1:
                continue

            headers = worksheet_headers[header_index]
            keys = [normalize_key(str(header or "")) for header in headers]
            has_rows = False
            for empty_row_count, values in enumerate(
                worksheet.iter_rows(min_row=header_index + 2, values_only=True),
                start=1,
            ):
                row = {
                    keys[index]: values[index]
                    for index in range(min(len(keys), len(values)))
                    if keys[index]
                }
                if any(value is not None and str(value or "").strip() != "" for value in row.values()):
                    has_rows = True
                    break
                if empty_row_count >= 50:
                    break

            if has_rows and best_score > best_sheet_score:
                selected_sheet = worksheet
                selected_header_index = header_index
                selected_headers = headers
                best_sheet_score = best_score

        if selected_sheet is None:
            return []

        keys = [normalize_key(str(header or "")) for header in selected_headers]
        rows: list[ReportRow] = []
        consecutive_empty = 0
        for values in selected_sheet.iter_rows(min_row=selected_header_index + 2, values_only=True):
            row = {keys[index]: values[index] for index in range(min(len(keys), len(values))) if keys[index]}
            if any(value is not None and str(value or "").strip() != "" for value in row.values()):
                rows.append(row)
                consecutive_empty = 0
            else:
                consecutive_empty += 1
                if consecutive_empty >= 50:
                    break

        return rows
    finally:
        workbook.close()
